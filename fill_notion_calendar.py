"""엑셀 노션_캘린더 시트를 노션 캘린더 DB에 채운다.

속성 매핑:
  이벤트명 → 이름(title)
  날짜/종료일 → 날짜(date)
  프로젝트 → 프로젝트명(select, 단일)
  유형     → 유형(select)
  시간     → 시간(rich_text)
  참석자   → 참석자(rich_text)
  장소     → 장소(rich_text)
  메모     → 페이지 본문
  No·이벤트ID → 제외

  python fill_notion_calendar.py --db <ID>              속성 준비 + 기존행 삭제 + 전체 삽입
  python fill_notion_calendar.py --db <ID> --limit 1    1건만
  python fill_notion_calendar.py --db <ID> --keep       기존 행 삭제 안 함
"""
from __future__ import annotations

import argparse
from pathlib import Path

import httpx
from openpyxl import load_workbook

KEY_NAME = "NOTION_PRIVATESPACE_API"
XLSX = r"C:\Users\마준서\Documents\카카오톡 받은 파일\디지털트윈_일정_교차검증_테스트데이터셋_1.xlsx"

# select로 만들 속성(정해진 값들). rich_text로 만들 속성(자유 텍스트).
SELECT_PROPS = ["프로젝트명", "유형"]
TEXT_PROPS = ["시간", "참석자", "장소"]


def _load_key() -> str:
    for line in (Path(__file__).parent / ".env").read_text(encoding="utf-8").splitlines():
        if "=" in line and line.split("=", 1)[0].strip().lower() == KEY_NAME.lower():
            return line.split("=", 1)[1].strip().strip("'\"")
    raise SystemExit(f"[X] .env에 {KEY_NAME}가 없습니다.")


BASE = "https://api.notion.com/v1"
HEADERS = {
    "Authorization": f"Bearer {_load_key()}",
    "Notion-Version": "2022-06-28",
    "Content-Type": "application/json",
}


def ensure_properties(db_id: str) -> None:
    """필요한 속성이 없으면 만든다. 기존 프로젝트명(multi_select)은 select로 바꾼다."""
    current = httpx.get(f"{BASE}/databases/{db_id}", headers=HEADERS, timeout=15).json()
    existing = current.get("properties", {})

    changes: dict = {}
    for name in SELECT_PROPS:
        # 이미 select면 그대로. 다른 타입(multi_select 등)이거나 없으면 select로.
        if existing.get(name, {}).get("type") != "select":
            changes[name] = {"select": {}}
    for name in TEXT_PROPS:
        if name not in existing:
            changes[name] = {"rich_text": {}}

    if not changes:
        print("속성 이미 준비됨.")
        return
    r = httpx.patch(
        f"{BASE}/databases/{db_id}", headers=HEADERS, json={"properties": changes}, timeout=15
    )
    if r.status_code == 200:
        print(f"속성 준비 완료: {list(changes)}")
    else:
        print(f"[X] 속성 준비 실패: {r.json().get('message', r.text)[:150]}")
        raise SystemExit(1)


def clear_rows(db_id: str) -> int:
    """DB의 기존 행(페이지)을 모두 보관처리(archive)한다."""
    removed = 0
    while True:
        q = httpx.post(f"{BASE}/databases/{db_id}/query", headers=HEADERS, timeout=15).json()
        results = q.get("results", [])
        if not results:
            break
        for page in results:
            httpx.patch(
                f"{BASE}/pages/{page['id']}", headers=HEADERS, json={"archived": True}, timeout=15
            )
            removed += 1
        if not q.get("has_more"):
            break
    return removed


def _val(row: dict, col: str) -> str:
    v = row.get(col)
    return "" if v in (None, "None") else str(v)


def add_event(db_id: str, row: dict) -> tuple[bool, str]:
    name = _val(row, "이벤트명") or "제목 없음"
    props: dict = {"이름": {"title": [{"text": {"content": name}}]}}

    if _val(row, "날짜"):
        d: dict = {"start": _val(row, "날짜")}
        if _val(row, "종료일"):
            d["end"] = _val(row, "종료일")
        props["날짜"] = {"date": d}

    if _val(row, "프로젝트"):
        props["프로젝트명"] = {"select": {"name": _val(row, "프로젝트")}}
    if _val(row, "유형"):
        props["유형"] = {"select": {"name": _val(row, "유형")}}
    for col in ("시간", "참석자", "장소"):
        if _val(row, col):
            props[col] = {"rich_text": [{"text": {"content": _val(row, col)}}]}

    children = []
    if _val(row, "메모"):
        children.append({
            "object": "block", "type": "paragraph",
            "paragraph": {"rich_text": [{"type": "text", "text": {"content": _val(row, "메모")}}]},
        })

    body = {"parent": {"database_id": db_id}, "properties": props, "children": children}
    r = httpx.post(f"{BASE}/pages", headers=HEADERS, json=body, timeout=20)
    if r.status_code == 200:
        return True, name
    return False, f"{name}: {r.json().get('message', r.text)[:120]}"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", required=True)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--keep", action="store_true", help="기존 행 삭제 안 함")
    args = parser.parse_args()

    db_id = args.db.split("?")[0].replace("-", "")

    ensure_properties(db_id)
    if not args.keep:
        n = clear_rows(db_id)
        print(f"기존 행 {n}개 정리")

    wb = load_workbook(XLSX, data_only=True, read_only=True)
    rows = list(wb["노션_캘린더"].iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    records = [dict(zip(header, r)) for r in rows[1:]]
    if args.limit:
        records = records[: args.limit]

    print(f"\n{len(records)}건 삽입 시작...\n")
    ok = 0
    for row in records:
        success, msg = add_event(db_id, row)
        if success:
            ok += 1
        else:
            print(f"  [X] {msg}")
    print(f"\n완료: {ok}/{len(records)}건 삽입")


if __name__ == "__main__":
    main()
